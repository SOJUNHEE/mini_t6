"""파이프라인 검증 — 샘플 엑셀로 실제 실행한다.

data/samples/axport_upload_sample_v0.1.xlsx 에 의도적으로 심은
오류 4종이 전부 잡히는지 확인한다.

업로드 제한값은 config/settings.toml 에서 미정("")이다. 테스트는
환경변수로 **시험용 값**을 주입한다. settings.toml 은 비워 둔다.
"""

from __future__ import annotations

import os
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SAMPLE = PROJECT_ROOT / "data" / "samples" / "axport_upload_sample_v0.1.xlsx"
TEMPLATE = PROJECT_ROOT / "data" / "samples" / "axport_upload_template_v0.1.xlsx"

# 승인 전 시험용 값. settings.toml 에 넣지 않는다.
TEST_ENV = {
    "AXPORT_FEATURES__UPLOAD": "true",
    "AXPORT_UPLOAD__ENABLED": "true",
    "AXPORT_UPLOAD__MAX_FILE_BYTES": "10485760",
    "AXPORT_UPLOAD__MAX_DECOMPRESSED_BYTES": "52428800",
    "AXPORT_UPLOAD__MAX_ROWS_PER_SHEET": "10000",
    "AXPORT_UPLOAD__MAX_SHEETS": "20",
    "AXPORT_UPLOAD__MAX_CELLS": "200000",
    "AXPORT_UPLOAD__PROCESSING_TIMEOUT_SECONDS": "30",
}


@pytest.fixture()
def ctx(tmp_path_factory):
    """테스트는 instance/uploads 를 건드리지 않고 임시 폴더를 쓴다."""
    from axport.config.settings import get_settings
    from axport.upload_validation import pipeline as P

    uploads = tmp_path_factory.mktemp("uploads")
    env = {**TEST_ENV, "AXPORT_STORAGE__UPLOADS_DIR": str(uploads)}
    for k, v in env.items():
        os.environ[k] = v
    get_settings.cache_clear()
    yield P.build_context(get_settings())
    for k in env:
        os.environ.pop(k, None)
    get_settings.cache_clear()


@pytest.fixture()
def uploaded(ctx):
    from axport.upload_validation import pipeline as P

    assert SAMPLE.is_file(), f"샘플 파일이 없습니다: {SAMPLE}"
    return ctx, P.ingest(ctx, SAMPLE.read_bytes(), SAMPLE.name)


# ── ① 업로드 ──────────────────────────────────────────────────────────
def test_limits_block_upload_when_undecided():
    """제한값이 미정이면 업로드를 거부한다. 임의 기본값을 쓰지 않는다."""
    from axport.config.settings import get_settings
    from axport.upload_validation import limits as L

    for k in TEST_ENV:
        os.environ.pop(k, None)
    get_settings.cache_clear()
    with pytest.raises(L.UploadRejected) as exc:
        L.load_limits(get_settings())
    assert exc.value.code == "limits_undecided"
    assert exc.value.detail["undecided"]


def test_server_generates_path_not_filename(uploaded):
    _, res = uploaded
    up = res["upload"]
    assert up["display_name"] == SAMPLE.name
    # 저장 경로에 원본 파일명이 들어가지 않는다
    assert SAMPLE.stem not in up["source_path"]
    assert up["source_path"].endswith("/source.xlsx")
    assert up["upload_id"] in up["source_path"]
    assert up["sha256"]


def test_rejects_non_xlsx(ctx):
    from axport.upload_validation import limits as L
    from axport.upload_validation import pipeline as P

    with pytest.raises(L.UploadRejected) as exc:
        P.ingest(ctx, b"not a zip file at all", "fake.xlsx")
    assert exc.value.code == "not_a_valid_xlsx"


def test_rejects_macro_extension(ctx):
    from axport.upload_validation import limits as L
    from axport.upload_validation import pipeline as P

    with pytest.raises(L.UploadRejected) as exc:
        P.ingest(ctx, SAMPLE.read_bytes(), "book.xlsm")
    assert exc.value.code == "macro_format_not_allowed"


def test_rejects_oversized(ctx):
    from axport.upload_validation import limits as L
    from axport.upload_validation import pipeline as P

    ctx.limits.max_file_bytes = 100
    with pytest.raises(L.UploadRejected) as exc:
        P.ingest(ctx, SAMPLE.read_bytes(), SAMPLE.name)
    assert exc.value.code == "file_too_large"


# ── ② 검증 — 의도한 오류 4종 ──────────────────────────────────────────
@pytest.fixture()
def validation(uploaded):
    from axport.upload_validation import pipeline as P

    ctx, res = uploaded
    return ctx, res, P.run_validation(ctx, res["upload"]["upload_id"])


def _find(report, code, sheet=None):
    return [f for f in report["findings"]
            if f["code"] == code and (sheet is None or f["sheet"] == sheet)]


def test_case_1a_real_zero_is_recorded_as_actual(validation):
    """오류 케이스 1 — 실제 값 0 을 미입력으로 처리하지 않는다."""
    _, _, report = validation
    zeros = _find(report, "value_is_real_zero")
    assert zeros, "값 0 을 '실제 값 0' 으로 기록하지 않았다"
    rows = {(f["sheet"], f["row"]) for f in zeros}
    assert ("수출실적", 5) in rows      # S-003 전량 취소
    assert ("재고·생산", 4) in rows     # IV-002 재고 실제 0


def test_case_1b_blank_is_reported_not_filled(validation):
    """오류 케이스 1 — 미입력을 0 으로 채우지 않고 보고한다."""
    _, _, report = validation
    blanks = _find(report, "area_required_missing")
    assert blanks, "미입력 항목을 보고하지 않았다"
    hits = {(f["sheet"], f["row"], f["column"]) for f in blanks}
    assert ("수출예정거래", 5, "단가") in hits        # T-003
    assert ("원가·비용", 5, "보험료") in hits         # CT-003
    assert ("물류", 5, "운임") in hits                # L-003
    assert ("제품정보", 6, "전략물자해당") in hits     # P-004


def test_case_2_broken_reference(validation):
    """오류 케이스 2 — 거래처 시트와 연결되지 않는 행."""
    _, _, report = validation
    broken = _find(report, "broken_reference")
    assert len(broken) == 1, f"연결 끊김이 1건이어야 한다: {broken}"
    f = broken[0]
    assert f["severity"] == "error"
    assert f["sheet"] == "수출예정거래"
    assert f["row"] == 6                       # T-004
    assert f["column"] == "거래처ID"
    assert f["detail"]["value"] == "BP-999"
    assert f["detail"]["target_sheet"] == "거래처"


def test_case_3_invalid_date(validation):
    """오류 케이스 3 — 날짜 형식 오류."""
    _, _, report = validation
    bad = _find(report, "invalid_date")
    assert len(bad) == 1, f"날짜 오류가 1건이어야 한다: {bad}"
    f = bad[0]
    assert f["severity"] == "error"
    assert (f["sheet"], f["row"], f["column"]) == ("수출실적", 7, "거래일")
    assert "2026/13/45" in f["message"]


def test_case_4_duplicate_row(validation):
    """오류 케이스 4 — 중복 행."""
    _, _, report = validation
    dup_rows = _find(report, "duplicate_row")
    dup_keys = _find(report, "duplicate_key")
    assert len(dup_rows) == 1, f"중복 행이 1건이어야 한다: {dup_rows}"
    assert dup_rows[0]["sheet"] == "수출실적"
    assert dup_rows[0]["row"] == 6
    assert dup_rows[0]["detail"]["first_row"] == 5
    # 주키 중복도 함께 잡힌다
    assert any(d["detail"]["value"] == "S-003" for d in dup_keys)


def test_validation_has_no_formula_cells(validation):
    _, _, report = validation
    assert not _find(report, "formula_cell")


def test_validation_counts_and_blocking(validation):
    _, _, report = validation
    assert report["counts"]["error"] > 0
    assert report["blocking"] is True
    for f in report["findings"]:
        assert f["severity"] in ("error", "warning", "info")
        if f["code"] not in ("sheet_missing", "sheet_unknown"):
            assert f["row"] is not None


# ── ③ 매핑 ────────────────────────────────────────────────────────────
def test_mapping_blocked_while_errors_remain(validation):
    from axport.upload_validation import pipeline as P

    ctx, res, _ = validation
    with pytest.raises(P.StageBlocked) as exc:
        P.run_mapping(ctx, res["upload"]["upload_id"])
    assert exc.value.code == "validation_has_errors"


def test_mapping_on_clean_template(ctx):
    """양식 파일은 오류가 없으므로 매핑까지 진행된다."""
    from axport.upload_validation import pipeline as P

    res = P.ingest(ctx, TEMPLATE.read_bytes(), TEMPLATE.name)
    uid = res["upload"]["upload_id"]
    v = P.run_validation(ctx, uid)
    assert v["counts"]["error"] == 0, v["findings"]
    m = P.run_mapping(ctx, uid)
    assert m["needs_confirmation_count"] == 0
    assert m["unmapped_count"] == 0
    assert all(x["status"] == "exact" for x in m["mappings"])


def test_ambiguous_columns_need_confirmation(ctx, tmp_path):
    """별칭·유사 컬럼명은 후보로만 제시하고 확정하지 않는다."""
    from openpyxl import load_workbook

    from axport.upload_validation import pipeline as P

    wb = load_workbook(TEMPLATE)
    ws = wb["제품정보"]
    ws.cell(row=1, column=5).value = "HS CODE"    # HS코드 → 별칭
    ws.cell(row=1, column=8).value = "원산지"      # 제조국코드 → 별칭
    altered = tmp_path / "altered.xlsx"
    wb.save(altered)

    res = P.ingest(ctx, altered.read_bytes(), "altered.xlsx")
    uid = res["upload"]["upload_id"]
    P.run_validation(ctx, uid)
    m = P.run_mapping(ctx, uid)

    assert m["needs_confirmation_count"] >= 2
    assert m["ready_for_analysis"] is False
    pending = {x["template_column"]: x for x in m["mappings"]
               if x["status"] == "needs_confirmation"}
    assert "HS코드" in pending
    assert pending["HS코드"]["candidates"][0]["uploaded_column"] == "HS CODE"
    assert pending["HS코드"]["uploaded_column"] is None  # 확정하지 않았다


def test_analysis_blocked_until_mapping_confirmed(ctx, tmp_path):
    from openpyxl import load_workbook

    from axport.upload_validation import pipeline as P

    wb = load_workbook(TEMPLATE)
    wb["제품정보"].cell(row=1, column=5).value = "HS CODE"
    altered = tmp_path / "altered2.xlsx"
    wb.save(altered)

    res = P.ingest(ctx, altered.read_bytes(), "altered2.xlsx")
    uid = res["upload"]["upload_id"]
    P.run_validation(ctx, uid)
    P.run_mapping(ctx, uid)

    with pytest.raises(P.StageBlocked) as exc:
        P.list_subjects(ctx, uid)
    assert exc.value.code == "mapping_needs_confirmation"

    # 사용자 확인 후에는 진행된다
    P.confirm_mapping(ctx, uid, {"제품정보.HS코드": "HS CODE"})
    subjects = P.list_subjects(ctx, uid)
    assert subjects["count"] == 0   # 양식은 데이터가 없다


# ── ④ 평가 대상 ───────────────────────────────────────────────────────
@pytest.fixture()
def analyzable(ctx, tmp_path):
    """샘플에서 오류 행을 뺀 사본으로 분석 단계까지 간다."""
    from openpyxl import load_workbook

    from axport.upload_validation import pipeline as P

    wb = load_workbook(SAMPLE)
    wb["수출실적"].delete_rows(6, 2)            # 중복 행 + 날짜 오류 행
    wb["수출예정거래"].cell(row=6, column=6).value = "BP-001"  # 연결 복구
    fixed = tmp_path / "fixed.xlsx"
    wb.save(fixed)

    res = P.ingest(ctx, fixed.read_bytes(), "fixed.xlsx")
    uid = res["upload"]["upload_id"]
    v = P.run_validation(ctx, uid)
    assert v["counts"]["error"] == 0, [f for f in v["findings"]
                                      if f["severity"] == "error"]
    P.run_mapping(ctx, uid)
    return ctx, uid


def test_subjects_use_evaluation_unit(analyzable):
    from axport.upload_validation import pipeline as P

    ctx, uid = analyzable
    out = P.list_subjects(ctx, uid)
    assert out["count"] == 4
    s = out["subjects"][0]
    for key in ("company_id", "product_id", "model_name",
                "destination_country_code", "evaluation_base_date",
                "trade_terms", "analysis_mode"):
        assert key in s


def test_market_exploration_when_trade_terms_missing(analyzable):
    """거래 조건이 없으면 시장 탐색으로 표시한다."""
    from axport.rules_engine import subject as S
    from axport.upload_validation import pipeline as P

    ctx, uid = analyzable
    out = P.list_subjects(ctx, uid)
    by_product = {s["product_id"]: s for s in out["subjects"]}

    t003 = by_product["P-002"]   # 단가·통화·금액배율 미입력
    assert t003["analysis_mode"] == S.MARKET_EXPLORATION
    assert t003["analysis_mode_label"] == "시장 탐색"
    assert "완결된 거래 판정이 아닙니다" in t003["disclaimer"]
    assert "단가" in t003["missing_trade_terms"]

    t001 = by_product["P-001"]   # 거래 조건 완비
    assert t001["analysis_mode"] == S.DEAL_EVALUATION
    assert "disclaimer" not in t001


# ── ⑤ 3층 결과 ────────────────────────────────────────────────────────
@pytest.fixture()
def analysis(analyzable):
    from axport.upload_validation import pipeline as P

    ctx, uid = analyzable
    subjects = P.list_subjects(ctx, uid)
    sid = next(s["subject_id"] for s in subjects["subjects"]
               if s["product_id"] == "P-001")
    return P.analyze(ctx, uid, sid)


def test_three_layers_present(analysis):
    r = analysis["result"]
    assert set(r) == {"mandatory_review", "business_fitness",
                      "evidence_status", "overall"}


def test_regulation_is_withheld_because_source_missing(analysis):
    m = analysis["result"]["mandatory_review"]
    assert m["status"] == "withheld"
    assert m["status_reason_code"] == "rule_source_missing"
    assert m["blocking_pending_count"] == 1
    assert any("별표2의2" in s
               for c in m["pending_checks"] for s in c["required_sources"])


def test_business_score_cannot_offset(analysis):
    """게이트 입력에 사업성이 들어가지 않는다."""
    o = analysis["result"]["overall"]
    assert o["business_fitness_is_gate_input"] is False
    assert o["gate_inputs"] == ["mandatory_review", "evidence_status"]
    assert o["decision"] == "withheld"       # G1 에서 막힌다
    assert not o["gate_trace"][0]["passed"]
    assert "상쇄하지 않는다" in o["disclaimer"]


def test_no_score_numbers_invented(analysis):
    """가중치·임계값 미승인이므로 점수가 나오지 않는다."""
    b = analysis["result"]["business_fitness"]
    assert b["composite"]["state"] == "not_evaluated"
    assert "value" not in b["composite"]
    assert b["composite"]["reason_code"] == "scoring_config_unapproved"
    for area in b["areas"]:
        assert area["subscore"]["state"] != "actual"
        assert "value" not in area["subscore"]
    assert b["scoring_config_ref"]["review_status"] == "미검토"


def test_unevaluable_areas_report_reason_and_required_inputs(analysis):
    """평가 불가 영역은 이유와 필요한 입력을 함께 반환한다."""
    for area in analysis["result"]["business_fitness"]["areas"]:
        assert area["subscore"]["reason_code"]
        assert area["required_to_evaluate"], area["area_code"]
        assert area["subscore"]["display_label"] in ("자료 부족", "평가 보류")


def test_evidence_status_reports_sources(analysis):
    e = analysis["result"]["evidence_status"]
    assert e["overall"] == "insufficient"
    assert e["data_sources"]
    verified = [s for s in e["data_sources"] if s["verified"] == "검증완료"]
    names = {s["api_name"] for s in verified}
    assert any("품목별 국가별 수출입실적" in n for n in names)
    assert e["rule_readiness"][0]["usable_in_production"] is False


def test_main_summary_has_five_items(analysis):
    """design.md — 메인은 5개 항목, 각 핵심 결과 1개 + 설명 1줄."""
    items = analysis["main_summary"]
    assert [i["item"] for i in items] == [
        "regulation", "market", "price", "logistics", "stability"]
    assert [i["label"] for i in items] == [
        "규제", "시장성", "가격", "물류", "안정성"]
    assert items[0]["layer"] == "mandatory_review"
    for i in items[1:]:
        assert i["layer"] == "business_fitness"
    for i in items:
        assert i["headline"]
        assert "\n" not in i["note"]

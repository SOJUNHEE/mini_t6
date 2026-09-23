"""공용 픽스처.

기본 검증은 네트워크·비밀키 없이 재현 가능하다 (CONTRIBUTING §12).
외부 API 응답은 api-vault/specs/samples 에 기록된 **실제 응답 구조**를
그대로 흉내낸 고정 XML 로 대체한다. 이 대체는 테스트에서만 일어나며,
런타임에는 장애 시 샘플로 대체하지 않는다 (명세서 §7).
"""

from __future__ import annotations

import os
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SAMPLE = PROJECT_ROOT / "data" / "samples" / "axport_upload_sample_v0.1.xlsx"
TEMPLATE = PROJECT_ROOT / "data" / "samples" / "axport_upload_template_v0.1.xlsx"

# 승인 전 시험용 제한값. config/settings.toml 에 넣지 않는다.
UPLOAD_ENV = {
    "AXPORT_FEATURES__UPLOAD": "true",
    "AXPORT_UPLOAD__ENABLED": "true",
    "AXPORT_UPLOAD__MAX_FILE_BYTES": "10485760",
    "AXPORT_UPLOAD__MAX_DECOMPRESSED_BYTES": "52428800",
    "AXPORT_UPLOAD__MAX_ROWS_PER_SHEET": "10000",
    "AXPORT_UPLOAD__MAX_SHEETS": "20",
    "AXPORT_UPLOAD__MAX_CELLS": "200000",
    "AXPORT_UPLOAD__PROCESSING_TIMEOUT_SECONDS": "30",
}

PERIOD = {"start": "2026-01", "end": "2026-03"}


@pytest.fixture()
def env(tmp_path_factory):
    """업로드·보고서·캐시를 임시 폴더로 격리한다."""
    from axport.config.settings import get_settings

    overrides = {
        **UPLOAD_ENV,
        "AXPORT_STORAGE__UPLOADS_DIR": str(tmp_path_factory.mktemp("uploads")),
        "AXPORT_STORAGE__REPORTS_DIR": str(tmp_path_factory.mktemp("reports")),
        "AXPORT_STORAGE__CACHE_DIR": str(tmp_path_factory.mktemp("cache")),
    }
    for k, v in overrides.items():
        os.environ[k] = v
    get_settings.cache_clear()
    yield overrides
    for k in overrides:
        os.environ.pop(k, None)
    get_settings.cache_clear()


@pytest.fixture()
def ctx(env):
    from axport.config.settings import get_settings
    from axport.upload_validation import pipeline as P

    return P.build_context(get_settings())


@pytest.fixture()
def clean_workbook(tmp_path):
    """샘플에서 의도된 오류 행을 뺀 사본. 분석 단계까지 진행 가능하다."""
    from openpyxl import load_workbook

    wb = load_workbook(SAMPLE)
    wb["수출실적"].delete_rows(6, 2)                  # 중복·날짜오류 행
    wb["수출예정거래"].cell(row=6, column=6).value = "BP-001"   # 연결 복구
    path = tmp_path / "clean.xlsx"
    wb.save(path)
    return path


# ── 외부 응답 고정본 ──────────────────────────────────────────────────
def canned_xml(rows: list[dict]) -> bytes:
    """관세청 nitemtrade 실제 응답 구조 (specs/samples 기록과 동일)."""
    items = "".join(
        "<item>"
        f"<balPayments>{r['bal']}</balPayments>"
        f"<expDlr>{r['exp']}</expDlr><expWgt>{r['expw']}</expWgt>"
        f"<hsCd>{r['hs']}</hsCd>"
        f"<impDlr>{r['imp']}</impDlr><impWgt>{r['impw']}</impWgt>"
        f"<statCd>{r['cc']}</statCd>"
        f"<statCdCntnKor1>{r['cn']}</statCdCntnKor1>"
        f"<statKor>{r['name']}</statKor><year>{r['period']}</year>"
        "</item>"
        for r in rows
    )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        "<response><header><resultCode>00</resultCode>"
        "<resultMsg>정상서비스.</resultMsg></header>"
        f"<body><items>{items}</items></body></response>"
    ).encode("utf-8")


CANNED_ROWS = [
    {"period": "총계", "hs": "-", "cc": "-", "cn": "-", "name": "-",
     "exp": "300", "expw": "30", "imp": "120", "impw": "12", "bal": "180"},
    {"period": "2026.01", "hs": "854232", "cc": "US", "cn": "미국",
     "name": "기억소자", "exp": "100", "expw": "10", "imp": "40",
     "impw": "4", "bal": "60"},
    {"period": "2026.02", "hs": "854232", "cc": "US", "cn": "미국",
     "name": "기억소자", "exp": "200", "expw": "20", "imp": "80",
     "impw": "8", "bal": "120"},
]

ERROR_XML = (
    "<OpenAPI_ServiceResponse><cmmMsgHeader>"
    "<errMsg>SERVICE ERROR</errMsg>"
    "<returnAuthMsg>SERVICE_KEY_IS_NOT_REGISTERED_ERROR</returnAuthMsg>"
    "<returnReasonCode>30</returnReasonCode>"
    "</cmmMsgHeader></OpenAPI_ServiceResponse>"
).encode("utf-8")

EMPTY_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    "<response><header><resultCode>00</resultCode>"
    "<resultMsg>정상서비스.</resultMsg></header>"
    "<body><items></items></body></response>"
).encode("utf-8")


@pytest.fixture()
def stub_customs(monkeypatch):
    """외부 호출을 고정 응답으로 바꾼다. 키도 필요 없게 만든다."""
    from axport.external_data.providers import customs_trade as ct
    from axport.rules_engine import evaluate as E

    state = {"body": canned_xml(CANNED_ROWS), "status": 200, "calls": []}

    def fake_get(url, timeout):
        # URL 에 키가 섞여 있어도 기록하지 않는다
        query = url.split("serviceKey=")[-1].split("&", 1)[-1]
        state["calls"].append(query)
        # 실측과 같게: 10자리 HS 는 0건, 6자리는 데이터 (customs_trade.py 주석)
        hs = next((kv.split("=")[1] for kv in query.split("&")
                   if kv.startswith("hsSgn=")), "")
        if len(hs) > 6 and state["body"] is not ERROR_XML:
            return state["status"], EMPTY_XML
        return state["status"], state["body"]

    def fake_key(*_a, **_k):
        return "TEST-KEY-NOT-REAL"

    monkeypatch.setattr(ct, "_http_get", fake_get)
    monkeypatch.setattr(E, "_service_key", fake_key, raising=False)
    return state

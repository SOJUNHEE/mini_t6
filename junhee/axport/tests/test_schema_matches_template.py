"""schema.py 선언이 실제 양식 파일과 일치하는지 확인한다.

둘이 어긋나면 검증이 조용히 잘못 동작한다. 여기서 잡는다.
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from axport.upload_validation import schema

TEMPLATE = (Path(__file__).resolve().parents[1]
            / "data" / "samples" / "axport_upload_template_v0.1.xlsx")

TAG_RE = re.compile(r"^\[(공통필수|영역필수|선택)(?::([^\]]+))?\]")


def test_template_file_exists():
    assert TEMPLATE.is_file(), f"양식 파일이 없습니다: {TEMPLATE}"


def test_sheet_names_match():
    wb = load_workbook(TEMPLATE, read_only=True)
    actual = [s for s in wb.sheetnames if s not in schema.NON_DATA_SHEETS]
    wb.close()
    assert actual == list(schema.SHEET_NAMES)


def test_template_version_matches():
    wb = load_workbook(TEMPLATE, read_only=True)
    ws = wb[schema.INFO_SHEET]
    version = None
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == "template_version":
            version = str(row[1]).strip()
            break
    wb.close()
    assert version == schema.TEMPLATE_VERSION


def test_columns_and_requirement_classes_match():
    wb = load_workbook(TEMPLATE, read_only=True)
    problems: list[str] = []
    for sheet_def in schema.SHEETS:
        ws = wb[sheet_def.name]
        rows = ws.iter_rows(values_only=True)
        headers = [("" if h is None else str(h).strip()) for h in next(rows)]
        while headers and headers[-1] == "":
            headers.pop()
        descs = [("" if d is None else str(d).strip()) for d in next(rows)]

        if headers != list(sheet_def.column_names):
            problems.append(
                f"{sheet_def.name}: 컬럼 불일치\n"
                f"  양식   {headers}\n  schema {list(sheet_def.column_names)}"
            )
            continue

        for name, desc in zip(headers, descs):
            m = TAG_RE.match(desc)
            if not m:
                problems.append(f"{sheet_def.name}.{name}: 2행 태그 없음 ({desc!r})")
                continue
            expected = schema.TAG_TO_CLASS[m.group(1)]
            col = sheet_def.column(name)
            if col.requirement != expected:
                problems.append(
                    f"{sheet_def.name}.{name}: 필수구분 불일치 "
                    f"양식={expected} schema={col.requirement}"
                )
            if expected == schema.AREA_REQUIRED and (m.group(2) or "") != col.area:
                problems.append(
                    f"{sheet_def.name}.{name}: 영역명 불일치 "
                    f"양식={m.group(2)!r} schema={col.area!r}"
                )
    wb.close()
    assert not problems, "\n".join(problems)

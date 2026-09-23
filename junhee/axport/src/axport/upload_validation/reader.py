"""엑셀 판독 — 수식을 실행하지 않는다 (명세서 §5).

openpyxl 을 `data_only=False` 로 연다. 이 모드는 수식을 **문자열 그대로**
돌려주고 계산하지 않는다. 수식 셀을 발견하면 값으로 쓰지 않고 오류로 보고한다.
(`data_only=True` 는 Excel 이 마지막에 캐시한 계산값을 쓰는 것이므로,
신뢰할 수 없는 값을 실제 값처럼 읽게 된다. 쓰지 않는다.)
"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from axport.upload_validation import schema
from axport.upload_validation.limits import Limits, UploadRejected

FORMULA_PREFIXES = ("=", "+=", "-=")


@dataclass
class Cell:
    """읽은 셀 하나. 빈칸과 0 을 구분해 담는다."""

    raw: object | None
    is_blank: bool
    is_formula: bool
    text: str | None = None

    @property
    def is_present(self) -> bool:
        return not self.is_blank


@dataclass
class SheetData:
    name: str
    headers: list[str]
    rows: list[dict[str, Cell]] = field(default_factory=list)
    row_numbers: list[int] = field(default_factory=list)
    formula_cells: list[dict] = field(default_factory=list)


@dataclass
class WorkbookData:
    sheets: dict[str, SheetData] = field(default_factory=dict)
    extra_sheets: list[str] = field(default_factory=list)
    missing_sheets: list[str] = field(default_factory=list)
    info_sheet_values: dict[str, str] = field(default_factory=dict)
    total_cells: int = 0


def _to_cell(value: object | None) -> Cell:
    if value is None:
        return Cell(raw=None, is_blank=True, is_formula=False)
    if isinstance(value, str):
        stripped = value.strip()
        if stripped == "":
            # 공백만 있는 셀은 미입력으로 본다. 0 으로 바꾸지 않는다.
            return Cell(raw=None, is_blank=True, is_formula=False)
        if stripped.startswith(FORMULA_PREFIXES):
            return Cell(raw=stripped, is_blank=False, is_formula=True, text=stripped)
        return Cell(raw=stripped, is_blank=False, is_formula=False, text=stripped)
    if isinstance(value, (dt.datetime, dt.date)):
        return Cell(raw=value, is_blank=False, is_formula=False,
                    text=value.strftime("%Y-%m-%d"))
    return Cell(raw=value, is_blank=False, is_formula=False, text=str(value))


def read(path: Path, limits: Limits) -> WorkbookData:
    """양식 시트를 읽는다. 제한 초과 시 UploadRejected."""
    wb = load_workbook(
        path,
        data_only=False,   # 수식을 계산하지 않는다
        read_only=True,
        keep_links=False,
    )
    try:
        if len(wb.sheetnames) > limits.max_sheets:
            raise UploadRejected(
                "too_many_sheets",
                "시트 수 제한을 초과했습니다.",
                {"sheets": len(wb.sheetnames), "max_sheets": limits.max_sheets},
            )

        data = WorkbookData()
        present = set(wb.sheetnames)
        data.missing_sheets = [n for n in schema.SHEET_NAMES if n not in present]
        data.extra_sheets = [
            n for n in wb.sheetnames
            if n not in schema.SHEET_NAMES and n not in schema.NON_DATA_SHEETS
        ]

        if schema.INFO_SHEET in present:
            ws = wb[schema.INFO_SHEET]
            for row in ws.iter_rows(values_only=True):
                if row and row[0] is not None:
                    key = str(row[0]).strip()
                    val = "" if len(row) < 2 or row[1] is None else str(row[1]).strip()
                    data.info_sheet_values.setdefault(key, val)

        for sheet_def in schema.SHEETS:
            if sheet_def.name not in present:
                continue
            data.sheets[sheet_def.name] = _read_sheet(
                wb[sheet_def.name], sheet_def, limits, data
            )
        return data
    finally:
        wb.close()


def _read_sheet(ws, sheet_def: schema.Sheet, limits: Limits,
                acc: WorkbookData) -> SheetData:
    rows = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        return SheetData(name=sheet_def.name, headers=[])

    headers = [("" if h is None else str(h).strip()) for h in header_row]
    # 뒤쪽 빈 헤더 제거
    while headers and headers[-1] == "":
        headers.pop()

    out = SheetData(name=sheet_def.name, headers=headers)

    try:
        next(rows)  # 2행 = 설명. 데이터가 아니다
    except StopIteration:
        return out

    data_row_no = schema.DATA_START_ROW
    for values in rows:
        if values is None or all(v is None or str(v).strip() == "" for v in values):
            data_row_no += 1
            continue

        if len(out.rows) >= limits.max_rows_per_sheet:
            raise UploadRejected(
                "too_many_rows",
                "시트 행 수 제한을 초과했습니다.",
                {"sheet": sheet_def.name,
                 "max_rows_per_sheet": limits.max_rows_per_sheet},
            )

        record: dict[str, Cell] = {}
        for idx, name in enumerate(headers):
            if not name:
                continue
            value = values[idx] if idx < len(values) else None
            cell = _to_cell(value)
            record[name] = cell
            acc.total_cells += 1
            if cell.is_formula:
                out.formula_cells.append({
                    "sheet": sheet_def.name,
                    "row": data_row_no,
                    "column": name,
                })

        if acc.total_cells > limits.max_cells:
            raise UploadRejected(
                "too_many_cells",
                "셀 수 제한을 초과했습니다.",
                {"cells": acc.total_cells, "max_cells": limits.max_cells},
            )

        out.rows.append(record)
        out.row_numbers.append(data_row_no)
        data_row_no += 1

    return out
